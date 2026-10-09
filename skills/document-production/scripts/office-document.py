import argparse
import json
import os
from pathlib import Path
import tempfile
from xml.sax.saxutils import escape


def load_source(filename):
    source = json.loads(Path(filename).read_text(encoding="utf-8"))
    if not isinstance(source, dict) or set(source) != {"title", "blocks"}:
        raise ValueError("expected title and blocks")
    if not isinstance(source["title"], str) or not source["title"].strip():
        raise ValueError("title must be nonempty text")
    if not isinstance(source["blocks"], list):
        raise ValueError("blocks must be a list")
    items = []
    for block in source["blocks"]:
        if not isinstance(block, dict) or set(block) != {"type", "items"} or block["type"] != "bullets":
            raise ValueError("only bullet blocks are supported")
        if not isinstance(block["items"], list) or any(not isinstance(item, str) or not item.strip() for item in block["items"]):
            raise ValueError("bullet items must be nonempty text")
        items.extend(block["items"])
    return source["title"], items


def render_docx(output, title, items, font):
    from docx import Document
    document = Document()
    document.add_heading(title, 0)
    for item in items:
        document.add_paragraph(item, style="List Bullet")
    document.save(output)
    actual = [paragraph.text for paragraph in Document(output).paragraphs]
    if actual != [title, *items]:
        raise ValueError("DOCX read-back differs from source")
    return {"paragraphs": len(actual)}


def render_pptx(output, title, items, font):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    if len(title) > 80 or any(len(item) > 160 or "\n" in item for item in items):
        raise ValueError("PPTX sample needs title <=80 and single-line bullets <=160 characters; edit layout for longer text")
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    groups = [items[index:index + 5] for index in range(0, len(items), 5)] or [[]]
    for group in groups:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        text = slide.placeholders[1].text_frame
        text.clear()
        for index, item in enumerate(group):
            paragraph = text.paragraphs[0] if index == 0 else text.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(24)
    deck.save(output)
    reopened = Presentation(output)
    actual = [paragraph.text for slide in reopened.slides for paragraph in slide.placeholders[1].text_frame.paragraphs if paragraph.text]
    if actual != items or any(slide.shapes.title.text != title for slide in reopened.slides):
        raise ValueError("PPTX read-back differs from source")
    return {"slides": len(reopened.slides)}


def render_xlsx(output, title, items, font):
    from openpyxl import Workbook, load_workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Notes"
    for row, value in enumerate([title, *items], start=1):
        cell = sheet.cell(row=row, column=1, value=value)
        cell.data_type = "s"
    sheet.column_dimensions["A"].width = 80
    sheet.freeze_panes = "A2"
    workbook.save(output)
    workbook.close()
    reopened = load_workbook(output)
    try:
        cells = list(reopened.active["A"])
        if [cell.value for cell in cells] != [title, *items] or any(cell.data_type != "s" for cell in cells):
            raise ValueError("XLSX read-back differs from literal source text")
        return {"rows": len(cells), "formulasEvaluated": False}
    finally:
        reopened.close()


def render_pdf(output, title, items, font):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from pypdf import PdfReader
    if font:
        pdfmetrics.registerFont(TTFont("DocumentFont", str(font)))
    elif not all(value.isascii() for value in [title, *items]):
        raise ValueError("non-ASCII PDF text requires --font with a licensed TrueType font covering those glyphs")
    styles = getSampleStyleSheet()
    if font:
        styles["Title"].fontName = "DocumentFont"
        styles["BodyText"].fontName = "DocumentFont"
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for item in items:
        story.extend([Paragraph(escape(item), styles["BodyText"]), Spacer(1, 8)])
    SimpleDocTemplate(str(output)).build(story)
    with PdfReader(output) as document:
        actual = " ".join(" ".join(page.extract_text().split()) for page in document.pages)
        expected = " ".join(" ".join(value.split()) for value in [title, *items])
        if actual != expected:
            raise ValueError("PDF extracted text differs from source")
        return {"pages": len(document.pages)}


def main():
    parser = argparse.ArgumentParser(description="Generate and reopen a small editable Office or PDF document from title/bullet JSON.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--format", required=True, choices=["docx", "pptx", "xlsx", "pdf"])
    parser.add_argument("--font", type=Path)
    args = parser.parse_args()
    temporary = None
    try:
        if args.output.suffix.lower() != f".{args.format}":
            raise ValueError("output extension must match format")
        title, items = load_source(args.input)
        descriptor, filename = tempfile.mkstemp(suffix=f".{args.format}", dir=args.output.parent)
        os.close(descriptor)
        temporary = Path(filename)
        renderers = {"docx": render_docx, "pptx": render_pptx, "xlsx": render_xlsx, "pdf": render_pdf}
        details = renderers[args.format](temporary, title, items, args.font)
        os.replace(temporary, args.output)
        print(json.dumps({"output": str(args.output), "format": args.format, "reopened": True, **details}))
        return 0
    except ImportError as error:
        print(json.dumps({"error": "missing-dependency", "module": error.name}))
        return 2
    except (ValueError, OSError, KeyError) as error:
        print(json.dumps({"error": "document-failed", "detail": str(error)}))
        return 1
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    raise SystemExit(main())
