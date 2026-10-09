import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/document-production/scripts/office-document.py"


class DocumentDepthTests(unittest.TestCase):
    def generate(self, directory, format_name, source):
        source_path = Path(directory) / "source.json"
        source_path.write_text(json.dumps(source), encoding="utf-8")
        output = Path(directory) / f"notes.{format_name}"
        result = subprocess.run(
            [sys.executable, "-I", str(SCRIPT), "--input", str(source_path), "--output", str(output), "--format", format_name],
            capture_output=True, text=True, timeout=30,
        )
        return result, output

    def test_docx_reopens_paragraph_heading_ordered_list_and_table(self):
        from docx import Document
        source = {
            "title": "Harbor notes",
            "blocks": [
                {"type": "heading", "text": "Berth plan"},
                {"type": "paragraph", "text": "Draft stays at 12.4 m."},
                {"type": "list", "ordered": True, "items": ["Sound the basin", "Log the tide"]},
                {"type": "bullets", "items": ["Keep one fact per bullet"]},
                {"type": "table", "columns": ["Berth", "Draft m"], "rows": [["North", "12.4"], ["South", "9.8"]]},
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            result, output = self.generate(directory, "docx", source)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            document = Document(output)
            paragraphs = [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs]
            self.assertEqual(paragraphs, [
                ("Title", "Harbor notes"),
                ("Heading 2", "Berth plan"),
                ("Normal", "Draft stays at 12.4 m."),
                ("List Number", "Sound the basin"),
                ("List Number", "Log the tide"),
                ("List Bullet", "Keep one fact per bullet"),
            ])
            self.assertEqual(len(document.tables), 1)
            table = document.tables[0]
            self.assertEqual([[cell.text for cell in row.cells] for row in table.rows], [
                ["Berth", "Draft m"], ["North", "12.4"], ["South", "9.8"],
            ])

    def test_xlsx_table_is_a_real_grid_with_literal_formula_text(self):
        from openpyxl import load_workbook
        source = {
            "title": "Harbor notes",
            "blocks": [
                {"type": "table", "columns": ["Berth", "Expr"], "rows": [["North", "=SUM(A1:A2)"], ["South", "9.8"]]},
            ],
        }
        with tempfile.TemporaryDirectory() as directory:
            result, output = self.generate(directory, "xlsx", source)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output)
            sheet = workbook.active
            self.assertEqual(sheet["A1"].value, "Harbor notes")
            self.assertEqual([[cell.value for cell in row] for row in sheet["A3:B5"]], [
                ["Berth", "Expr"], ["North", "=SUM(A1:A2)"], ["South", "9.8"],
            ])
            self.assertTrue(all(cell.data_type == "s" for row in sheet["A3:B5"] for cell in row))
            workbook.close()

    def test_old_bullets_schema_still_renders(self):
        source = {"title": "Harbor notes", "blocks": [{"type": "bullets", "items": ["Keep one fact per bullet"]}]}
        for format_name in ["docx", "xlsx"]:
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                result, _ = self.generate(directory, format_name, source)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_malformed_table_is_rejected_and_preserves_existing_output(self):
        for format_name in ["docx", "xlsx"]:
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / f"notes.{format_name}"
                output.write_bytes(b"existing document")
                result, _ = self.generate(directory, format_name, {
                    "title": "Notes", "blocks": [{"type": "table", "columns": ["A", "B"], "rows": [["only-one"]]}],
                })
                self.assertEqual(result.returncode, 1)
                self.assertIn("row 1 has 1 cells but 2 columns", result.stdout)
                self.assertEqual(output.read_bytes(), b"existing document")

    def test_pptx_and_pdf_reject_tables_explicitly(self):
        source = {"title": "Notes", "blocks": [{"type": "table", "columns": ["A"], "rows": [["1"]]}]}
        for format_name in ["pptx", "pdf"]:
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                result, output = self.generate(directory, format_name, source)
                self.assertEqual(result.returncode, 1)
                self.assertIn(f"{format_name} does not support table blocks", result.stdout)
                self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
