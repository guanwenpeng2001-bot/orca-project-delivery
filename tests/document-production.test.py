import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "skills/document-production/scripts/office-document.py"


class DocumentProductionTests(unittest.TestCase):
    def generate(self, directory, format_name, source):
        source_path = Path(directory) / "source.json"
        source_path.write_text(json.dumps(source), encoding="utf-8")
        output = Path(directory) / f"notes.{format_name}"
        result = subprocess.run([sys.executable, "-I", str(SCRIPT), "--input", str(source_path), "--output", str(output), "--format", format_name], capture_output=True, text=True, timeout=30)
        return result, output

    def test_all_formats_generate_and_reopen_source_text(self):
        for format_name in ["docx", "pptx", "xlsx", "pdf"]:
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                source = {"title": "Harbor notes", "blocks": [{"type": "bullets", "items": ["Keep one fact per bullet", "Amount: 125.50 USD", "Harbor <north> & south"]}]}
                result, output = self.generate(directory, format_name, source)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                self.assertTrue(json.loads(result.stdout)["reopened"])
                self.assertTrue(output.read_bytes().startswith(b"%PDF" if format_name == "pdf" else b"PK"))
                if format_name == "docx":
                    from docx import Document
                    self.assertEqual(Document(output).paragraphs[2].text, "Amount: 125.50 USD")
                if format_name == "pptx":
                    from pptx import Presentation
                    self.assertEqual(Presentation(output).slides[0].placeholders[1].text_frame.paragraphs[1].text, "Amount: 125.50 USD")
                if format_name == "xlsx":
                    from openpyxl import load_workbook
                    workbook = load_workbook(output)
                    self.assertEqual(workbook.active["A3"].value, "Amount: 125.50 USD")
                    workbook.close()
                if format_name == "pdf":
                    from pypdf import PdfReader
                    with PdfReader(output) as document:
                        self.assertIn("Amount: 125.50 USD", document.pages[0].extract_text())

    def test_unsupported_content_does_not_replace_existing_document(self):
        for format_name in ["docx", "pptx", "xlsx", "pdf"]:
            with self.subTest(format=format_name), tempfile.TemporaryDirectory() as directory:
                output = Path(directory) / f"notes.{format_name}"
                output.write_bytes(b"existing document")
                result, _ = self.generate(directory, format_name, {"title": "Notes", "blocks": [{"type": "table", "items": []}]})
                self.assertEqual(result.returncode, 1)
                self.assertEqual(output.read_bytes(), b"existing document")
                self.assertEqual(len(list(Path(directory).iterdir())), 2)

    def test_spreadsheet_formula_like_text_stays_literal(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            result, output = self.generate(directory, "xlsx", {"title": "=1+1", "blocks": [{"type": "bullets", "items": ["=SUM(A1:A2)"]}]})
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output)
            self.assertEqual(workbook.active["A1"].data_type, "s")
            self.assertEqual(workbook.active["A2"].value, "=SUM(A1:A2)")
            workbook.close()

    def test_pdf_without_unicode_font_fails_without_artifact(self):
        with tempfile.TemporaryDirectory() as directory:
            result, output = self.generate(directory, "pdf", {"title": "港口", "blocks": []})
            self.assertEqual(result.returncode, 1)
            self.assertFalse(output.exists())

    def test_slide_overflow_is_paginated_and_overlong_text_rejected(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            result, output = self.generate(directory, "pptx", {"title": "Notes", "blocks": [{"type": "bullets", "items": [f"Item {number}" for number in range(12)]}]})
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertEqual(len(Presentation(output).slides), 3)
            original = output.read_bytes()
            result, _ = self.generate(directory, "pptx", {"title": "Notes", "blocks": [{"type": "bullets", "items": ["long " * 100]}]})
            self.assertEqual(result.returncode, 1)
            self.assertEqual(output.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
